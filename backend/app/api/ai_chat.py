from fastapi import APIRouter, Depends, HTTPException
from sqlalchemy.orm import Session
from typing import Dict, Any, Optional, List
from datetime import datetime, date, timedelta
import logging
import json
import re

from app.core.database import get_db
from app.models.order import Order, OrderStatus
from app.models.client import Client
from app.models.ai_chat_history import AIChatHistory
from app.schemas.order import OrderCreate
from app.services.ai_chat_service import AIChatService

logger = logging.getLogger(__name__)

router = APIRouter()


@router.post("/process")
async def process_chat_message(
    payload: Dict[str, Any],
    db: Session = Depends(get_db)
):
    """
    AI 채팅 메시지 처리 및 주문 등록/수정
    
    Args:
        payload: {
            "message": str,  # 사용자 입력 메시지
            "context": {
                "pending_order": dict,  # 확인 대기 중인 주문
                "recent_messages": list,  # 최근 대화 기록
                "confirm": bool  # 주문 확인 여부
            }
        }
    
    Returns:
        {
            "intent": str,  # "create_order", "update_order", "confirm_order", "order_created", etc.
            "message": str,  # AI 응답 메시지
            "parsed_order": dict,  # 추출된 주문 정보
            "action": str,  # "confirm_order", "create_order", etc.
            "order_created": dict  # 생성된 주문 정보
        }
    """
    
    try:
        message = payload.get("message", "").strip()
        context = payload.get("context", {})
        pending_order = context.get("pending_order")
        confirm = context.get("confirm", False)
        recent_messages = context.get("recent_messages", [])
        model_name = payload.get("model", "auto")  # AI 모델 선택
        
        if not message:
            raise HTTPException(status_code=400, detail="메시지가 비어있습니다.")
        
        # AI 서비스 초기화 (선택된 모델 사용)
        ai_service = AIChatService(model_name=model_name)
        
        # 확인 요청인 경우 주문 생성
        pending_orders = context.get("pending_orders")
        if confirm and (pending_order or pending_orders):
            try:
                # 여러 주문 생성 (pending_orders가 리스트인 경우)
                if pending_orders and isinstance(pending_orders, list):
                    created_orders = []
                    for order_data in pending_orders:
                        order = await _create_order_from_parsed_data(db, order_data)
                        created_orders.append(order)
                    
                    order_numbers = [o.order_number for o in created_orders]
                    return {
                        "intent": "orders_created",
                        "message": f"✅ {len(created_orders)}개 주문이 등록되었습니다!\n\n주문번호:\n" + "\n".join([f"• {num}" for num in order_numbers]),
                        "parsed_order": None,
                        "action": None,
                        "orders_created": [{"order_number": o.order_number, "id": o.id} for o in created_orders]
                    }
                else:
                    # 단일 주문 생성
                    order = await _create_order_from_parsed_data(db, pending_order)
                    return {
                        "intent": "order_created",
                        "message": f"✅ 주문이 등록되었습니다!\n\n주문번호: {order.order_number}\n온도대: {order.temperature_zone}\n팔레트: {order.pallet_count}개\n중량: {order.weight_kg}kg",
                        "parsed_order": None,
                        "action": None,
                        "order_created": {
                            "order_number": order.order_number,
                            "id": order.id
                        }
                    }
            except Exception as e:
                logger.error(f"주문 생성 실패: {e}")
                return {
                    "intent": "error",
                    "message": f"❌ 주문 등록에 실패했습니다: {str(e)}",
                    "parsed_order": None,
                    "action": None,
                    "order_created": None
                }
        
        # AI로 메시지 파싱
        result = await ai_service.process_message(
            message=message,
            context={
                "pending_order": pending_order,
                "recent_messages": recent_messages
            },
            db=db
        )
        
        intent = result.get("intent", "unknown")
        parsed_order = result.get("parsed_order")
        parsed_orders = result.get("parsed_orders")  # 여러 주문
        ai_message = result.get("message", "처리할 수 없는 요청입니다.")
        
        # 📝 히스토리 저장
        try:
            history = AIChatHistory(
                user_id=context.get("user_id"),  # 나중에 사용자 관리 추가 시
                session_id=context.get("session_id"),
                user_message=message,
                assistant_message=ai_message,
                intent=intent,
                action=result.get("action"),
                parsed_order=parsed_order,
                parsed_orders=parsed_orders,
                dispatch_recommendation=result.get("dispatch_recommendation")
            )
            db.add(history)
            db.commit()
            logger.info(f"✅ 채팅 히스토리 저장: {history.id}")
        except Exception as e:
            logger.error(f"❌ 히스토리 저장 실패: {e}")
            # 히스토리 저장 실패는 전체 요청을 실패시키지 않음
            db.rollback()
        
        # 여러 주문 생성 의도
        if intent == "create_multiple_orders" and parsed_orders:
            return {
                "intent": "confirm_multiple_orders",
                "message": ai_message,
                "parsed_orders": parsed_orders,
                "parsed_order": None,
                "action": "confirm_multiple_orders",
                "order_created": None
            }
        
        # 주문 수정 의도
        if intent == "update_order" and parsed_order:
            order_number = parsed_order.get("order_number")
            if order_number:
                try:
                    order = db.query(Order).filter(Order.order_number == order_number).first()
                    if not order:
                        return {
                            "intent": "error",
                            "message": f"❌ 주문번호 {order_number}을(를) 찾을 수 없습니다.",
                            "parsed_order": None,
                            "action": None,
                            "order_created": None
                        }
                    
                    # 주문 업데이트
                    for key, value in parsed_order.items():
                        if key != "order_number" and value is not None:
                            setattr(order, key, value)
                    
                    db.commit()
                    db.refresh(order)
                    
                    return {
                        "intent": "order_updated",
                        "message": f"✅ 주문 {order_number}이(가) 수정되었습니다!",
                        "parsed_order": None,
                        "action": None,
                        "order_created": None
                    }
                except Exception as e:
                    logger.error(f"주문 수정 실패: {e}")
                    db.rollback()
                    return {
                        "intent": "error",
                        "message": f"❌ 주문 수정에 실패했습니다: {str(e)}",
                        "parsed_order": None,
                        "action": None,
                        "order_created": None
                    }
        
        # 주문 생성 의도 - 확인 요청
        if intent == "create_order" and parsed_order:
            return {
                "intent": "confirm_order",
                "message": ai_message,
                "parsed_order": parsed_order,
                "action": "confirm_order",
                "order_created": None
            }
        
        # 일반 응답
        return {
            "intent": intent,
            "message": ai_message,
            "parsed_order": parsed_order,
            "action": None,
            "order_created": None
        }
        
    except Exception as e:
        logger.error(f"AI 채팅 처리 오류: {e}")
        raise HTTPException(status_code=500, detail=f"처리 중 오류가 발생했습니다: {str(e)}")


async def _create_order_from_parsed_data(db: Session, parsed_order: Dict[str, Any]) -> Order:
    """
    파싱된 데이터로 주문 생성
    """
    try:
        # 주문번호 생성
        if not parsed_order.get("order_number"):
            timestamp = int(datetime.now().timestamp() * 1000)
            parsed_order["order_number"] = f"ORD-{timestamp}"
        
        # 주문 날짜 설정
        if not parsed_order.get("order_date"):
            parsed_order["order_date"] = date.today()
        
        # 상태 설정
        parsed_order["status"] = OrderStatus.PENDING
        
        # Order 객체 생성
        order = Order(**parsed_order)
        
        db.add(order)
        db.commit()
        db.refresh(order)
        
        logger.info(f"✅ 주문 생성 완료: {order.order_number}")
        return order
        
    except Exception as e:
        db.rollback()
        logger.error(f"주문 생성 실패: {e}")
        raise e


@router.get("/history")
async def get_chat_history(
    limit: int = 50,
    offset: int = 0,
    intent: Optional[str] = None,
    session_id: Optional[str] = None,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    db: Session = Depends(get_db)
):
    """
    AI 채팅 히스토리 조회
    
    Args:
        limit: 조회할 개수 (기본 50)
        offset: 건너뛸 개수 (페이징용)
        intent: 의도 필터 (create_order, update_order 등)
        session_id: 세션 ID 필터
        start_date: 시작 날짜 (YYYY-MM-DD)
        end_date: 종료 날짜 (YYYY-MM-DD)
    
    Returns:
        {
            "total": int,
            "items": [
                {
                    "id": int,
                    "user_message": str,
                    "assistant_message": str,
                    "intent": str,
                    "created_at": str,
                    ...
                }
            ]
        }
    """
    try:
        query = db.query(AIChatHistory)
        
        # 필터 적용
        if intent:
            query = query.filter(AIChatHistory.intent == intent)
        
        if session_id:
            query = query.filter(AIChatHistory.session_id == session_id)
        
        if start_date:
            start_dt = datetime.strptime(start_date, "%Y-%m-%d")
            query = query.filter(AIChatHistory.created_at >= start_dt)
        
        if end_date:
            end_dt = datetime.strptime(end_date, "%Y-%m-%d") + timedelta(days=1)
            query = query.filter(AIChatHistory.created_at < end_dt)
        
        # 전체 개수
        total = query.count()
        
        # 정렬 및 페이징
        histories = query.order_by(AIChatHistory.created_at.desc()).offset(offset).limit(limit).all()
        
        # 결과 변환
        items = []
        for h in histories:
            items.append({
                "id": h.id,
                "user_id": h.user_id,
                "session_id": h.session_id,
                "user_message": h.user_message,
                "assistant_message": h.assistant_message,
                "intent": h.intent,
                "action": h.action,
                "parsed_order": h.parsed_order,
                "parsed_orders": h.parsed_orders,
                "dispatch_recommendation": h.dispatch_recommendation,
                "created_at": h.created_at.isoformat() if h.created_at else None
            })
        
        return {
            "total": total,
            "items": items
        }
        
    except Exception as e:
        logger.error(f"히스토리 조회 오류: {e}")
        raise HTTPException(status_code=500, detail=f"히스토리 조회 중 오류가 발생했습니다: {str(e)}")


@router.delete("/history/all")
async def delete_all_chat_history(
    db: Session = Depends(get_db)
):
    """
    모든 채팅 히스토리 삭제
    """
    try:
        deleted_count = db.query(AIChatHistory).delete()
        db.commit()
        
        return {
            "message": f"{deleted_count}개의 히스토리가 삭제되었습니다.",
            "deleted_count": deleted_count
        }
        
    except Exception as e:
        logger.error(f"전체 히스토리 삭제 오류: {e}")
        db.rollback()
        raise HTTPException(status_code=500, detail=f"히스토리 삭제 중 오류가 발생했습니다: {str(e)}")


@router.delete("/history/{history_id}")
async def delete_chat_history(
    history_id: int,
    db: Session = Depends(get_db)
):
    """
    특정 채팅 히스토리 삭제
    """
    try:
        history = db.query(AIChatHistory).filter(AIChatHistory.id == history_id).first()
        
        if not history:
            raise HTTPException(status_code=404, detail="히스토리를 찾을 수 없습니다.")
        
        db.delete(history)
        db.commit()
        
        return {"message": "히스토리가 삭제되었습니다."}
        
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"히스토리 삭제 오류: {e}")
        db.rollback()
        raise HTTPException(status_code=500, detail=f"히스토리 삭제 중 오류가 발생했습니다: {str(e)}")


@router.get("/history/export/excel")
async def export_chat_history_excel(
    db: Session = Depends(get_db)
):
    """
    채팅 히스토리를 Excel 파일로 내보내기
    """
    try:
        from fastapi.responses import StreamingResponse
        import io
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        
        # 히스토리 조회
        histories = db.query(AIChatHistory).order_by(AIChatHistory.created_at.desc()).all()
        
        # Excel 워크북 생성
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Chat History"
        
        # 헤더 스타일
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        header_alignment = Alignment(horizontal="center", vertical="center")
        
        # 헤더 작성
        headers = ["ID", "사용자 메시지", "AI 응답", "의도", "액션", "생성일시"]
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
        
        # 데이터 작성
        for row_idx, history in enumerate(histories, 2):
            ws.cell(row=row_idx, column=1, value=history.id)
            ws.cell(row=row_idx, column=2, value=history.user_message)
            ws.cell(row=row_idx, column=3, value=history.assistant_message)
            ws.cell(row=row_idx, column=4, value=history.intent or "")
            ws.cell(row=row_idx, column=5, value=history.action or "")
            ws.cell(row=row_idx, column=6, value=history.created_at.strftime("%Y-%m-%d %H:%M:%S") if history.created_at else "")
        
        # 열 너비 조정
        ws.column_dimensions['A'].width = 10
        ws.column_dimensions['B'].width = 50
        ws.column_dimensions['C'].width = 50
        ws.column_dimensions['D'].width = 20
        ws.column_dimensions['E'].width = 20
        ws.column_dimensions['F'].width = 20
        
        # 파일 저장
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        filename = f"chat_history_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )
        
    except Exception as e:
        logger.error(f"Excel 내보내기 오류: {e}")
        raise HTTPException(status_code=500, detail=f"Excel 내보내기 중 오류가 발생했습니다: {str(e)}")


@router.get("/history/stats")
async def get_chat_history_stats(
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    db: Session = Depends(get_db)
):
    """
    채팅 히스토리 통계
    
    Returns:
        {
            "total_conversations": int,
            "by_intent": {
                "create_order": int,
                "create_multiple_orders": int,
                ...
            },
            "recent_7_days": [
                {"date": "2026-02-01", "count": 10},
                ...
            ]
        }
    """
    try:
        query = db.query(AIChatHistory)
        
        # 날짜 필터
        if start_date:
            start_dt = datetime.strptime(start_date, "%Y-%m-%d")
            query = query.filter(AIChatHistory.created_at >= start_dt)
        
        if end_date:
            end_dt = datetime.strptime(end_date, "%Y-%m-%d") + timedelta(days=1)
            query = query.filter(AIChatHistory.created_at < end_dt)
        
        # 전체 대화 수
        total_conversations = query.count()
        
        # 의도별 통계
        by_intent = {}
        all_histories = query.all()
        for h in all_histories:
            intent = h.intent or "unknown"
            by_intent[intent] = by_intent.get(intent, 0) + 1
        
        # 최근 7일 통계
        recent_7_days = []
        for i in range(7):
            target_date = date.today() - timedelta(days=i)
            count = db.query(AIChatHistory).filter(
                AIChatHistory.created_at >= datetime.combine(target_date, datetime.min.time()),
                AIChatHistory.created_at < datetime.combine(target_date + timedelta(days=1), datetime.min.time())
            ).count()
            recent_7_days.append({
                "date": target_date.isoformat(),
                "count": count
            })
        
        return {
            "total_conversations": total_conversations,
            "by_intent": by_intent,
            "recent_7_days": recent_7_days
        }
        
    except Exception as e:
        logger.error(f"통계 조회 오류: {e}")
        raise HTTPException(status_code=500, detail=f"통계 조회 중 오류가 발생했습니다: {str(e)}")
